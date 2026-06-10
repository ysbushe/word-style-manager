"""Feature tests for style selection, reports, task schemes, and template editing."""

from __future__ import annotations

import json
import os
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from lxml import etree

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from src.ewt.core.editor import clean_old_template_versions, save_template_edits
from src.ewt.core.numbering import _clean_unused_numbering_root, _deduplicate_numbering_root
from src.ewt.core.updater import (
    _safe_extract_zip,
    _select_zip_asset,
    install_downloaded_update,
    is_newer_version,
)
from src.ewt.config import W_NS
from src.ewt.utils.display import chinese_style_name
from src.ewt.core.style_engine import (
    _clean_unused_styles_root,
    _hide_unused_builtin_styles_root,
    _reachable_style_and_numbering_ids,
    inspect_document,
    process_styles,
)
from src.ewt.ui.app import merge_document_paths
from src.ewt.ui.template_editor import (
    INDENT_CHOICES,
    SIZE_CHOICES,
    STYLE_VALUE_OPTIONS,
    _format_number,
    _raw_value,
)
from src.ewt.core.templates import (
    delete_template_from_library,
    export_style_template,
    import_styles_to_document,
    list_template_library,
    list_task_schemes,
    save_task_scheme,
)
from src.ewt.utils.safe_io import atomic_write_text, preserve_corrupt_file
from src.ewt.utils.paths import (
    config_path,
    is_legacy_bundled_template_dir,
    resolve_template_library_dir,
)
from src.ewt.utils.word_io import _copy_docx_with_replacements


class FeatureTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.source = self.root / "source.docx"
        self.target = self.root / "target.docx"
        doc = Document()
        used = doc.styles.add_style("工程标题", WD_STYLE_TYPE.PARAGRAPH)
        used.font.name = "Microsoft YaHei"
        doc.styles.add_style("显式保留但未使用", WD_STYLE_TYPE.PARAGRAPH)
        doc.add_paragraph("标题", style="工程标题")
        doc.save(self.source)
        target = Document()
        target.add_paragraph("正文")
        target.save(self.target)

    def tearDown(self):
        self.temp.cleanup()

    def test_selected_unused_style_can_be_exported(self):
        info = inspect_document(str(self.source))
        selected = next(item["style_id"] for item in info["styles"] if item["name"] == "显式保留但未使用")
        result = export_style_template(
            str(self.source),
            str(self.root),
            "选择模板",
            ["dotx"],
            str(self.root / "library"),
            [selected],
            True,
        )
        self.assertTrue(result["success"])
        exported = inspect_document(result["outputs"]["dotx"])
        self.assertIn("显式保留但未使用", {item["name"] for item in exported["styles"]})

    def test_export_can_keep_unused_styles_when_auto_clean_is_disabled(self):
        result = export_style_template(
            str(self.source),
            str(self.root),
            "未清理模板",
            ["dotx"],
            clean_unused=False,
        )
        self.assertTrue(result["success"])
        exported = inspect_document(result["outputs"]["dotx"])
        self.assertIn("显式保留但未使用", {item["name"] for item in exported["styles"]})
        self.assertEqual(result["removed_unused"], 0)

    def test_library_export_reuses_dotx_in_same_directory(self):
        result = export_style_template(
            str(self.source),
            str(self.root),
            "模板库复用",
            ["dotx", "library"],
            str(self.root),
        )
        self.assertTrue(result["success"])
        self.assertEqual(result["outputs"]["dotx"], result["outputs"]["library"])
        self.assertEqual(len(list(self.root.glob("模板库复用*.dotx"))), 1)
        index = (self.root / "template_index.json").read_text(encoding="utf-8")
        self.assertNotIn(str(self.root), index)

    def test_template_library_prunes_deleted_and_migrates_old_paths(self):
        library = self.root / "library"
        library.mkdir()
        existing = library / "现有模板.dotx"
        existing.write_bytes(self.source.read_bytes())
        stale = library / "已删除模板.dotx"
        old_root = self.root / "old-dist" / "templates"
        records = [
            {
                "name": "现有模板",
                "path": str(old_root / existing.name),
                "created_at": "2026-06-09 12:00:00",
            },
            {
                "name": "已删除模板",
                "path": str(stale),
                "created_at": "2026-06-09 11:00:00",
            },
        ]
        (library / "template_index.json").write_text(
            json.dumps(records, ensure_ascii=False),
            encoding="utf-8",
        )

        items = list_template_library(library)

        self.assertEqual([item["name"] for item in items], ["现有模板"])
        self.assertEqual(Path(items[0]["path"]), existing.resolve())
        repaired = json.loads((library / "template_index.json").read_text(encoding="utf-8"))
        self.assertEqual(repaired[0]["path"], existing.name)
        self.assertNotIn("已删除模板", {item["name"] for item in repaired})

    def test_template_library_delete_removes_file_and_index_record(self):
        library = self.root / "delete-library"
        library.mkdir()
        template = library / "待删除模板.dotx"
        template.write_bytes(self.source.read_bytes())
        (library / "template_index.json").write_text(
            json.dumps(
                [{"name": "待删除模板", "path": template.name, "created_at": "2026-06-10 10:00:00"}],
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )

        removed = delete_template_from_library(library, template)

        self.assertEqual(Path(removed), template.resolve())
        self.assertFalse(template.exists())
        self.assertEqual(
            json.loads((library / "template_index.json").read_text(encoding="utf-8")),
            [],
        )

    def test_config_path_is_anchored_to_application_directory(self):
        application_dir = self.root / "portable"
        self.assertEqual(
            config_path(application_dir),
            application_dir / "cleaner_config.json",
        )

    def test_legacy_bundled_library_migrates_to_documents(self):
        application_dir = self.root / "Word 样式管理器"
        legacy = application_dir / "templates"
        legacy.mkdir(parents=True)
        template = legacy / "旧模板.dotx"
        template.write_bytes(self.source.read_bytes())
        documents_library = self.root / "Documents" / "模板库"

        with mock.patch(
            "src.ewt.utils.paths.default_template_library_dir",
            return_value=documents_library,
        ):
            resolved, migrated = resolve_template_library_dir(legacy, application_dir)

        self.assertTrue(is_legacy_bundled_template_dir(legacy, application_dir))
        self.assertTrue(migrated)
        self.assertEqual(resolved, documents_library)
        self.assertTrue((documents_library / template.name).is_file())

    def test_custom_template_library_is_preserved(self):
        application_dir = self.root / "Word 样式管理器"
        custom = self.root / "自定义模板目录"
        with mock.patch(
            "src.ewt.utils.paths.default_template_library_dir",
            return_value=self.root / "Documents" / "模板库",
        ):
            resolved, migrated = resolve_template_library_dir(custom, application_dir)
        self.assertFalse(migrated)
        self.assertEqual(resolved, custom)

    def test_template_export_strips_large_document_assets(self):
        large_source = self.root / "large-source.docx"
        large_source.write_bytes(self.source.read_bytes())
        with zipfile.ZipFile(large_source, "a", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("word/media/large-unused.bin", os.urandom(2_000_000))
        result = export_style_template(
            str(large_source),
            str(self.root),
            "轻量模板",
            ["dotx"],
        )
        self.assertTrue(result["success"])
        output = Path(result["outputs"]["dotx"])
        self.assertLess(output.stat().st_size, large_source.stat().st_size // 4)
        with zipfile.ZipFile(output, "r") as archive:
            self.assertFalse(any(name.startswith("word/media/") for name in archive.namelist()))
            self.assertIn("word/styles.xml", archive.namelist())

    def test_builtin_style_names_are_displayed_in_chinese(self):
        self.assertEqual(chinese_style_name("Normal", "Normal"), "正文")
        self.assertEqual(chinese_style_name("Heading 1", "Heading1"), "一级标题")
        info = inspect_document(str(self.source))
        normal = next(item for item in info["styles"] if item["style_id"] == "Normal")
        self.assertEqual(normal["display_name"], "正文")

    def test_template_editor_unit_labels_convert_to_ooxml_values(self):
        self.assertEqual(_raw_value("小四（12 磅）", SIZE_CHOICES), "12")
        self.assertEqual(_raw_value("1.27 厘米", INDENT_CHOICES, 567), "720")
        self.assertTrue(all("磅" in label or "号" in label for label in SIZE_CHOICES))
        self.assertTrue(all("厘米" in label for label in INDENT_CHOICES))

    def test_import_creates_html_report(self):
        info = inspect_document(str(self.source))
        selected = next(item["style_id"] for item in info["styles"] if item["name"] == "工程标题")
        exported = export_style_template(
            str(self.source), str(self.root), "报告模板", ["dotx"], None, [selected], True
        )
        report = self.root / "report.html"
        result = import_styles_to_document(
            exported["outputs"]["dotx"],
            str(self.target),
            str(self.root / "output.docx"),
            selected_styles=[selected],
            report_path=str(report),
        )
        self.assertTrue(result["success"])
        self.assertTrue(report.exists())
        self.assertIn("Word 样式导入报告", report.read_text(encoding="utf-8"))

    def test_template_edit_creates_version(self):
        info = inspect_document(str(self.source))
        selected = next(item["style_id"] for item in info["styles"] if item["name"] == "工程标题")
        exported = export_style_template(str(self.source), str(self.root), "编辑模板", ["dotx"])
        result = save_template_edits(
            exported["outputs"]["dotx"],
            {selected: {"name": "工程标题新版", "font": "SimSun", "size": "16", "bold": True}},
            [{"format": "第%1章", "numFmt": "decimal", "suffix": "space"}],
            True,
            True,
            False,
        )
        self.assertTrue(Path(result["output"]).exists())
        self.assertNotEqual(Path(result["output"]), Path(exported["outputs"]["dotx"]))

    def test_template_edit_writes_line_spacing_rule(self):
        info = inspect_document(str(self.source))
        selected = next(item["style_id"] for item in info["styles"] if item["name"] == "工程标题")
        exported = export_style_template(str(self.source), str(self.root), "行距模板", ["dotx"])
        result = save_template_edits(
            exported["outputs"]["dotx"],
            {selected: {"line": "360", "line_rule": "exact"}},
        )
        with zipfile.ZipFile(result["output"], "r") as archive:
            styles = etree.fromstring(archive.read("word/styles.xml"))
        style = next(
            node
            for node in styles.findall("w:style", {"w": W_NS})
            if node.get(f"{{{W_NS}}}styleId") == selected
        )
        spacing = style.find("w:pPr/w:spacing", {"w": W_NS})
        self.assertEqual(spacing.get(f"{{{W_NS}}}line"), "360")
        self.assertEqual(spacing.get(f"{{{W_NS}}}lineRule"), "exact")

    def test_numbering_preview_formats_common_number_styles(self):
        self.assertEqual(_format_number(3, "decimalZero"), "03")
        self.assertEqual(_format_number(2, "upperLetter"), "B")
        self.assertEqual(_format_number(4, "lowerRoman"), "iv")
        self.assertEqual(_format_number(12, "chineseCounting"), "十二")

    def test_style_measurement_dropdowns_follow_selected_units(self):
        self.assertIn("1.27", STYLE_VALUE_OPTIONS[("left", "厘米")])
        self.assertIn("12.7", STYLE_VALUE_OPTIONS[("left", "毫米")])
        self.assertIn("1.5", STYLE_VALUE_OPTIONS[("line", "倍")])
        self.assertIn("18", STYLE_VALUE_OPTIONS[("line", "磅")])

    def test_template_version_cleanup_never_deletes_original(self):
        library = self.root / "versions"
        library.mkdir()
        original = library / "工程模板.dotx"
        original.write_bytes(self.source.read_bytes())
        versions = []
        for index in range(4):
            version = library / f"工程模板_v2026060{index + 1}_120000.dotx"
            version.write_bytes(self.source.read_bytes())
            os.utime(version, (index + 1, index + 1))
            versions.append(version)

        preview = clean_old_template_versions(library, "工程模板", keep=2, dry_run=True)
        removed = clean_old_template_versions(library, "工程模板", keep=2)

        self.assertEqual(preview, removed)
        self.assertTrue(original.exists())
        self.assertEqual(len(list(library.glob("工程模板_v*.dotx"))), 2)

    def test_docx_output_replaces_target_only_after_successful_write(self):
        output = self.root / "atomic-output.docx"
        output.write_bytes(b"old-content")
        with mock.patch("zipfile.ZipFile.writestr", side_effect=OSError("disk full")):
            with self.assertRaises(OSError):
                _copy_docx_with_replacements(self.source, output, {})
        self.assertEqual(output.read_bytes(), b"old-content")
        self.assertFalse(any(self.root.glob(f".{output.name}.*.tmp")))

    def test_atomic_text_write_and_corrupt_backup(self):
        path = self.root / "settings.json"
        atomic_write_text(path, '{"ok": true}')
        self.assertEqual(path.read_text(encoding="utf-8"), '{"ok": true}')
        backup = preserve_corrupt_file(path)
        self.assertFalse(path.exists())
        self.assertEqual(backup.read_text(encoding="utf-8"), '{"ok": true}')

    def test_task_scheme_roundtrip(self):
        save_task_scheme(self.root, "每日任务", {"source": "a.dotx", "targets": ["b.docx"]})
        schemes = list_task_schemes(self.root)
        self.assertEqual(schemes[0]["name"], "每日任务")
        self.assertEqual(schemes[0]["source"], "a.dotx")

    def test_duplicate_abstract_numbering_is_merged(self):
        root = etree.Element(f"{{{W_NS}}}numbering", nsmap={"w": W_NS})
        for abs_id in ("1", "2"):
            abstract = etree.SubElement(root, f"{{{W_NS}}}abstractNum")
            abstract.set(f"{{{W_NS}}}abstractNumId", abs_id)
            level = etree.SubElement(abstract, f"{{{W_NS}}}lvl")
            level.set(f"{{{W_NS}}}ilvl", "0")
            fmt = etree.SubElement(level, f"{{{W_NS}}}numFmt")
            fmt.set(f"{{{W_NS}}}val", "decimal")
        num = etree.SubElement(root, f"{{{W_NS}}}num")
        num.set(f"{{{W_NS}}}numId", "9")
        ref = etree.SubElement(num, f"{{{W_NS}}}abstractNumId")
        ref.set(f"{{{W_NS}}}val", "2")
        removed = _deduplicate_numbering_root(root)
        self.assertEqual(removed, 1)
        self.assertEqual(ref.get(f"{{{W_NS}}}val"), "1")

    def test_unused_single_and_multilevel_lists_can_be_cleaned_separately(self):
        root = etree.Element(f"{{{W_NS}}}numbering", nsmap={"w": W_NS})
        for abs_id, levels in (("1", 1), ("2", 3)):
            abstract = etree.SubElement(root, f"{{{W_NS}}}abstractNum")
            abstract.set(f"{{{W_NS}}}abstractNumId", abs_id)
            for level in range(levels):
                lvl = etree.SubElement(abstract, f"{{{W_NS}}}lvl")
                lvl.set(f"{{{W_NS}}}ilvl", str(level))
            num = etree.SubElement(root, f"{{{W_NS}}}num")
            num.set(f"{{{W_NS}}}numId", abs_id)
            ref = etree.SubElement(num, f"{{{W_NS}}}abstractNumId")
            ref.set(f"{{{W_NS}}}val", abs_id)

        removed_nums, removed_abs = _clean_unused_numbering_root(
            root,
            set(),
            remove_single_level=False,
            remove_multilevel=True,
        )
        self.assertEqual((removed_nums, removed_abs), (1, 1))
        self.assertEqual(
            [node.get(f"{{{W_NS}}}numId") for node in root.findall("w:num", {"w": W_NS})],
            ["1"],
        )

    def test_update_version_comparison(self):
        self.assertTrue(is_newer_version("v0.2.1", "0.2.0"))
        self.assertFalse(is_newer_version("v0.2.0", "0.2.0"))
        self.assertFalse(is_newer_version("v0.1.9", "0.2.0"))

    def test_portable_release_asset_is_preferred(self):
        assets = [
            {"name": "source.zip", "browser_download_url": "source"},
            {"name": "Word样式管理器_v0.3.0_绿色版.zip", "browser_download_url": "portable"},
        ]
        self.assertEqual(_select_zip_asset(assets)["browser_download_url"], "portable")

    def test_update_zip_rejects_path_traversal(self):
        archive = self.root / "unsafe.zip"
        extract_dir = self.root / "extract"
        extract_dir.mkdir()
        with zipfile.ZipFile(archive, "w") as zf:
            zf.writestr("../outside.txt", "unsafe")
        with self.assertRaises(RuntimeError):
            _safe_extract_zip(archive, extract_dir)
        self.assertFalse((self.root / "outside.txt").exists())

    def test_update_script_contains_backup_and_rollback(self):
        download_root = self.root / "download"
        payload = download_root / "payload"
        app_dir = self.root / "app"
        payload.mkdir(parents=True)
        app_dir.mkdir()
        (payload / "Word 样式管理器.exe").write_bytes(b"new")
        executable = app_dir / "Word 样式管理器.exe"
        executable.write_bytes(b"old")

        with (
            mock.patch("src.ewt.core.updater.sys.frozen", True, create=True),
            mock.patch("src.ewt.core.updater.sys.executable", str(executable)),
            mock.patch("src.ewt.core.updater.subprocess.Popen") as popen,
        ):
            install_downloaded_update(download_root, payload)

        script = (download_root / "apply_update.ps1").read_text(encoding="utf-8-sig")
        self.assertIn("backup_", script)
        self.assertIn("robocopy $backup $target /E", script)
        self.assertIn("$existing.ContainsKey", script)
        self.assertIn("update_error.txt", script)
        popen.assert_called_once()

    def test_document_paths_are_shared_and_deduplicated(self):
        first = str(self.root / "first.docx")
        duplicate = str(self.root / "." / "first.docx")
        second = str(self.root / "second.doc")
        result = merge_document_paths([first], [duplicate, second], ["not-a-document.dotx"])
        self.assertEqual(result, [first, second])

    def test_inspected_styles_put_unused_first_then_group_by_type(self):
        info = inspect_document(str(self.source))
        order = {"paragraph": 0, "character": 1, "table": 2, "numbering": 3}
        keys = [
            (0 if item.get("count", 0) == 0 else 1, order.get(item.get("type"), 9))
            for item in info["styles"]
        ]
        self.assertEqual(keys, sorted(keys))

    def test_dependency_style_is_required_instead_of_cleanable(self):
        path = self.root / "dependency.docx"
        doc = Document()
        base = doc.styles.add_style("基础段落样式", WD_STYLE_TYPE.PARAGRAPH)
        child = doc.styles.add_style("使用中的派生样式", WD_STYLE_TYPE.PARAGRAPH)
        child.base_style = base
        doc.add_paragraph("使用派生样式", style=child)
        doc.save(path)

        info = inspect_document(str(path))
        base_info = next(item for item in info["styles"] if item["name"] == "基础段落样式")

        self.assertEqual(base_info["count"], 0)
        self.assertTrue(base_info["dependency_required"])
        self.assertFalse(base_info["cleanable"])

    def test_reachable_cleanup_breaks_unused_style_numbering_cycle(self):
        styles = etree.Element(f"{{{W_NS}}}styles", nsmap={"w": W_NS})
        for style_id in ("Normal", "UsedByList", "DeadCycle"):
            style = etree.SubElement(styles, f"{{{W_NS}}}style")
            style.set(f"{{{W_NS}}}styleId", style_id)
            style.set(f"{{{W_NS}}}type", "paragraph")
            if style_id != "Normal":
                style.set(f"{{{W_NS}}}customStyle", "1")
                p_pr = etree.SubElement(style, f"{{{W_NS}}}pPr")
                num_pr = etree.SubElement(p_pr, f"{{{W_NS}}}numPr")
                num_id = etree.SubElement(num_pr, f"{{{W_NS}}}numId")
                num_id.set(f"{{{W_NS}}}val", "1" if style_id == "UsedByList" else "2")

        numbering = etree.Element(f"{{{W_NS}}}numbering", nsmap={"w": W_NS})
        for value, style_id in (("1", "UsedByList"), ("2", "DeadCycle")):
            abstract = etree.SubElement(numbering, f"{{{W_NS}}}abstractNum")
            abstract.set(f"{{{W_NS}}}abstractNumId", value)
            level = etree.SubElement(abstract, f"{{{W_NS}}}lvl")
            level.set(f"{{{W_NS}}}ilvl", "0")
            p_style = etree.SubElement(level, f"{{{W_NS}}}pStyle")
            p_style.set(f"{{{W_NS}}}val", style_id)
            num = etree.SubElement(numbering, f"{{{W_NS}}}num")
            num.set(f"{{{W_NS}}}numId", value)
            ref = etree.SubElement(num, f"{{{W_NS}}}abstractNumId")
            ref.set(f"{{{W_NS}}}val", value)

        style_ids, num_ids = _reachable_style_and_numbering_ids(
            styles,
            numbering,
            set(),
            {"1"},
            {f"{{{W_NS}}}basedOn"},
        )

        self.assertIn("UsedByList", style_ids)
        self.assertNotIn("DeadCycle", style_ids)
        self.assertEqual(num_ids, {"1"})

    def test_default_table_style_is_marked_as_required(self):
        with zipfile.ZipFile(self.source, "r") as archive:
            styles = etree.fromstring(archive.read("word/styles.xml"))
            settings = etree.fromstring(archive.read("word/settings.xml"))
        table_style = etree.SubElement(styles, f"{{{W_NS}}}style")
        table_style.set(f"{{{W_NS}}}styleId", "RequiredTableStyle")
        table_style.set(f"{{{W_NS}}}type", "table")
        table_style.set(f"{{{W_NS}}}customStyle", "1")
        default_style = etree.SubElement(settings, f"{{{W_NS}}}defaultTableStyle")
        default_style.set(f"{{{W_NS}}}val", "RequiredTableStyle")
        path = self.root / "default-table.docx"
        _copy_docx_with_replacements(
            self.source,
            path,
            {
                "word/styles.xml": etree.tostring(styles, xml_declaration=True, encoding="UTF-8"),
                "word/settings.xml": etree.tostring(settings, xml_declaration=True, encoding="UTF-8"),
            },
        )

        info = inspect_document(str(path))
        item = next(row for row in info["styles"] if row["style_id"] == "RequiredTableStyle")
        self.assertTrue(item["dependency_required"])
        self.assertFalse(item["cleanable"])

    def test_cleaning_never_overwrites_source_or_unverified_existing_output(self):
        same_path = process_styles(str(self.source), str(self.source))
        self.assertFalse(same_path["success"])

        output = self.root / "existing.docx"
        output.write_bytes(b"existing-output")
        with mock.patch(
            "src.ewt.core.style_engine._validate_clean_output",
            side_effect=RuntimeError("verification failed"),
        ):
            result = process_styles(str(self.source), str(output))
        self.assertFalse(result["success"])
        self.assertEqual(output.read_bytes(), b"existing-output")
        self.assertFalse(any(self.root.glob(f".{output.name}.*.verify")))

    def test_deep_clean_synchronizes_styles_with_effects(self):
        with zipfile.ZipFile(self.source, "r") as archive:
            effects = etree.fromstring(archive.read("word/styles.xml"))
        extra = etree.SubElement(effects, f"{{{W_NS}}}style")
        extra.set(f"{{{W_NS}}}styleId", "EffectsOnlyUnused")
        extra.set(f"{{{W_NS}}}type", "paragraph")
        extra.set(f"{{{W_NS}}}customStyle", "1")
        source = self.root / "effects-source.docx"
        _copy_docx_with_replacements(
            self.source,
            source,
            {
                "word/stylesWithEffects.xml": etree.tostring(
                    effects,
                    xml_declaration=True,
                    encoding="UTF-8",
                )
            },
        )
        output = self.root / "effects-cleaned.docx"

        result = process_styles(
            str(source),
            str(output),
            options={"mode": "deep"},
        )

        self.assertTrue(result["success"])
        self.assertTrue(result["verified"])
        with zipfile.ZipFile(output, "r") as archive:
            cleaned_effects = etree.fromstring(archive.read("word/stylesWithEffects.xml"))
        ids = {
            node.get(f"{{{W_NS}}}styleId")
            for node in cleaned_effects.findall("w:style", {"w": W_NS})
        }
        self.assertNotIn("EffectsOnlyUnused", ids)

    def test_unused_builtin_styles_can_be_hidden_instead_of_removed(self):
        root = etree.Element(f"{{{W_NS}}}styles", nsmap={"w": W_NS})
        builtin = etree.SubElement(root, f"{{{W_NS}}}style")
        builtin.set(f"{{{W_NS}}}styleId", "Heading1")
        builtin.set(f"{{{W_NS}}}type", "paragraph")
        etree.SubElement(builtin, f"{{{W_NS}}}qFormat")
        custom = etree.SubElement(root, f"{{{W_NS}}}style")
        custom.set(f"{{{W_NS}}}styleId", "CustomUnused")
        custom.set(f"{{{W_NS}}}type", "paragraph")
        custom.set(f"{{{W_NS}}}customStyle", "1")

        hidden = _hide_unused_builtin_styles_root(root, set())
        removed = _clean_unused_styles_root(root, set(), preserve_builtin=True)

        self.assertEqual(hidden, 1)
        self.assertEqual(removed, 1)
        self.assertIsNotNone(builtin.find("w:semiHidden", {"w": W_NS}))
        self.assertIsNotNone(builtin.find("w:unhideWhenUsed", {"w": W_NS}))
        self.assertIsNone(builtin.find("w:qFormat", {"w": W_NS}))
        self.assertEqual(
            [node.get(f"{{{W_NS}}}styleId") for node in root.findall("w:style", {"w": W_NS})],
            ["Heading1"],
        )


if __name__ == "__main__":
    unittest.main()
